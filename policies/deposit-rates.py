"""Prints the banks' deposit rates in `crates/core/src/short_term/deposits.rs`,
from the Bank of Israel's figures: each bank's average rate on the unlinked
shekel deposits households opened at a fixed rate in a month, by term (קו
המשווה, the Excel behind its page), and its average variable rate on the
deposits for a day (the deposit that can be taken out any day, `any_day_kind`).

    uv run --with openpyxl python policies/deposit-rates.py [YYYY-MM]

Without a month, the latest one published. A rate of 0 is one the Bank of
Israel left out: it publishes none for a bank's term with fewer than ten
deposits in the month, or under 5% of the month's deposits. Of the fixed
rates, the daily and weekly terms are left out: a fixed deposit is kept for
whole months. Of the variable ones (sheet L7.6.2a), only the daily term is
read: what a deposit that can be taken out any day pays.
"""

import io
import sys
import urllib.request

import openpyxl

EXCEL = "https://www.boi.org.il/boi_files/Pikuah/g060a.xls"
FIXED = "L7.6.1a"
VARIABLE = "L7.6.2a"
DAILY = "יומי"
# The terms kept, in the Excel's words, and the names the Rust code gives them.
TERMS = {
    "עד חודש": "UpToAMonth",
    "חודש עד 3 חודשים": "UpToThreeMonths",
    "3 חודשים עד 6 חודשים": "UpToSixMonths",
    "6 חודשים עד שנה": "UpToAYear",
    "שנה עד שנתיים": "UpToTwoYears",
    "שנתיים עד 3 שנים": "UpToThreeYears",
    "3 שנים עד 5 שנים": "UpToFiveYears",
}


def download() -> openpyxl.Workbook:
    # The file is an .xlsx, whatever its name says.
    request = urllib.request.Request(EXCEL, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(request, timeout=120) as response:
        return openpyxl.load_workbook(io.BytesIO(response.read()), data_only=True)


def any_day(workbook: openpyxl.Workbook, month: str) -> None:
    """The variable rate on deposits for a day, each bank's, for `month`."""
    rows = [list(row) for row in workbook[VARIABLE].iter_rows(values_only=True)]
    names = next(index for index, row in enumerate(rows) if "Hapoalim" in row)
    column = next(
        (
            column
            for column, title in enumerate(rows[names - 2])
            if isinstance(title, str) and title.strip().endswith("הממוצעת " + DAILY)
        ),
        None,
    )
    if column is None:
        sys.exit("found no daily term in the variable rates: the Excel has changed")
    months = {row[1]: row for row in rows[names + 1 :] if isinstance(row[1], str) and row[1][:2] == "20"}
    row = months[month]
    print()
    print(f"Deposits for a day at a variable rate, opened in {month}, the average yearly rate (% a year):")
    for offset in range(11):
        name = rows[names][column + offset]
        rate = row[column + offset]
        # 0 is a rate the Bank of Israel left out; the floating noise of the
        # sheet's own sums is rounded away.
        print(f"{name:20}{'—' if not rate else round(rate, 2):>16}")


def main() -> None:
    workbook = download()
    sheet = workbook[FIXED]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    # The banks' English names are in the row that has "Hapoalim"; each
    # block of banks has its term two rows above its first column.
    names = next(index for index, row in enumerate(rows) if "Hapoalim" in row)
    blocks = {}
    for column, title in enumerate(rows[names - 2]):
        if isinstance(title, str) and "הממוצעת" in title:
            term = title.split("הממוצעת", 1)[1].strip()
            if term in TERMS:
                blocks[TERMS[term]] = column
    if len(blocks) != len(TERMS):
        sys.exit(f"found the terms {sorted(blocks)}: the Excel has changed")
    first = min(blocks.values())
    banks = [
        (column - first, name)
        for column, name in enumerate(rows[names])
        if isinstance(name, str) and first <= column < first + 11
    ]

    months = {row[1]: row for row in rows[names + 1 :] if isinstance(row[1], str) and row[1][:2] == "20"}
    month = sys.argv[1] if len(sys.argv) > 1 else max(months)
    row = months[month]
    print(f"Fixed-rate deposits opened in {month}, the average yearly rate by term (% a year):")
    print(f"{'':20}" + "".join(f"{term:>16}" for term in blocks))
    for offset, name in banks:
        rates = [row[blocks[term] + offset] for term in blocks]
        # 0 is a rate the Bank of Israel left out.
        shown = "".join(f"{'—' if not rate else round(rate, 2):>16}" for rate in rates)
        print(f"{name:20}{shown}")
    any_day(workbook, month)


if __name__ == "__main__":
    main()
